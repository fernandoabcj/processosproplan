"""Geração dos relatórios: painel HTML (arquivo único) e planilha Excel."""

from __future__ import annotations

import html
from datetime import datetime
from pathlib import Path

from .analise import AnaliseProcesso, Consolidado
from .config import Config
from .util import duracao


def _n(x: float, casas: int = 1) -> str:
    return f"{x:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _d(dt: datetime | None, hora: bool = False) -> str:
    if not dt:
        return "—"
    return dt.strftime("%d/%m/%Y %H:%M" if hora and (dt.hour or dt.minute) else "%d/%m/%Y")


def _e(x) -> str:
    return html.escape(str(x if x is not None else ""))


# ====================================================================== HTML

_CSS = """
:root{--bg:#f6f7f9;--card:#fff;--tx:#1d2330;--mut:#5d6677;--ln:#e2e5ea;--ac:#1f5fae;
--ac2:#e8f0fb;--warn:#b54708;--warnbg:#fff4e5;--bad:#b42318;--badbg:#fdecea;--ok:#067647;}
@media (prefers-color-scheme:dark){:root{--bg:#14171c;--card:#1d2128;--tx:#e6e8eb;--mut:#9aa3b2;
--ln:#2e343e;--ac:#7fb0ef;--ac2:#1f2b3d;--warn:#f5b26b;--warnbg:#3a2a17;--bad:#f28b82;--badbg:#3b1d1b;--ok:#6fcf97;}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--tx);
font:14px/1.45 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif}
header{padding:24px 16px 8px;max-width:1280px;margin:auto}
h1{font-size:22px;margin:0 0 4px}h2{font-size:17px;margin:0 0 12px}
.sub{color:var(--mut)}main{max-width:1280px;margin:auto;padding:0 16px 40px}
section{background:var(--card);border:1px solid var(--ln);border-radius:10px;padding:16px;margin:16px 0}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px}
.kpi{background:var(--card);border:1px solid var(--ln);border-radius:10px;padding:14px}
.kpi b{display:block;font-size:26px;font-variant-numeric:tabular-nums}.kpi span{color:var(--mut)}
.tw{overflow-x:auto}table{border-collapse:collapse;width:100%;font-variant-numeric:tabular-nums}
th,td{padding:6px 8px;border-bottom:1px solid var(--ln);text-align:left;vertical-align:top}
th{font-weight:600;color:var(--mut);font-size:12px;text-transform:uppercase;letter-spacing:.02em;
position:sticky;top:0;background:var(--card);cursor:pointer}
td.n,th.n{text-align:right}tr.pp td{background:var(--ac2)}
.tag{display:inline-block;padding:1px 7px;border-radius:999px;background:var(--ac2);color:var(--ac);
font-size:12px;font-weight:600;white-space:nowrap}
.al{color:var(--bad);background:var(--badbg)}.wn{color:var(--warn);background:var(--warnbg)}
.bar{height:8px;background:var(--ac);border-radius:4px;min-width:2px}
details{border-top:1px solid var(--ln);padding:8px 0}summary{cursor:pointer}
input[type=search]{width:100%;max-width:420px;padding:8px 10px;border:1px solid var(--ln);
border-radius:8px;background:var(--bg);color:var(--tx);margin-bottom:10px}
.grid2{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:16px}
small{color:var(--mut)}
"""

_JS = """
document.querySelectorAll('input[data-filtra]').forEach(function(inp){
  inp.addEventListener('input',function(){
    var q=inp.value.toLowerCase();
    document.querySelectorAll(inp.dataset.filtra).forEach(function(el){
      el.style.display=el.textContent.toLowerCase().indexOf(q)>=0?'':'none';});});});
document.querySelectorAll('th').forEach(function(th){th.addEventListener('click',function(){
  var tb=th.closest('table').tBodies[0];if(!tb)return;var i=[].indexOf.call(th.parentNode.children,th);
  var asc=th.dataset.asc!=='1';th.dataset.asc=asc?'1':'0';
  var rows=[].slice.call(tb.rows);rows.sort(function(a,b){
    var x=a.cells[i].dataset.v||a.cells[i].textContent,y=b.cells[i].dataset.v||b.cells[i].textContent;
    var nx=parseFloat(x),ny=parseFloat(y);
    var r=(!isNaN(nx)&&!isNaN(ny))?nx-ny:x.localeCompare(y,'pt-BR');return asc?r:-r;});
  rows.forEach(function(r){tb.appendChild(r);});});});
"""


def _tag_setor(sigla: str | None) -> str:
    return f'<span class="tag">{_e(sigla)}</span>' if sigla else ""


def _tabela_setores(c: Consolidado, cfg: Config) -> str:
    maior = max((s.media for s in c.setores), default=0) or 1
    linhas = []
    for s in c.setores:
        largura = 100 * s.media / maior
        acima = (f'<span class="tag al">{s.estoque_acima_limite}</span>'
                 if s.estoque_acima_limite else "0")
        linhas.append(
            f"<tr><td><b>{_e(s.sigla)}</b><br><small>{_e(s.nome)}</small></td>"
            f'<td class="n">{s.processos}</td><td class="n">{s.passagens}</td>'
            f'<td class="n" data-v="{s.media:.3f}">{duracao(s.media)}</td>'
            f'<td style="width:140px"><div class="bar" style="width:{largura:.0f}%"></div></td>'
            f'<td class="n" data-v="{s.mediana:.3f}">{duracao(s.mediana)}</td>'
            f'<td class="n" data-v="{s.maximo:.3f}">{duracao(s.maximo)}</td>'
            f'<td class="n" data-v="{s.media_aguardando:.3f}">{duracao(s.media_aguardando)}</td>'
            f'<td class="n">{s.em_estoque}</td>'
            f'<td class="n" data-v="{s.idade_media_estoque:.3f}">{duracao(s.idade_media_estoque)}</td>'
            f'<td class="n">{acima}</td></tr>'
        )
    return (
        '<div class="tw"><table><thead><tr><th>Setor</th><th class="n">Processos</th>'
        '<th class="n">Passagens</th><th class="n">Tempo médio por passagem</th><th></th>'
        '<th class="n">Mediana</th><th class="n">Máximo</th><th class="n">Média até receber</th>'
        f'<th class="n">Em estoque hoje</th><th class="n">Idade média estoque</th>'
        f'<th class="n">Estoque ≥ {cfg.dias_parado_no_setor} dias</th></tr></thead><tbody>'
        + "".join(linhas)
        + "</tbody></table></div><p><small>Média, mediana e máximo consideram apenas passagens "
        "concluídas (o processo já saiu do setor). Tempo corrido.</small></p>"
    )


def _tabela_estoque(analises: list[AnaliseProcesso]) -> str:
    atuais = sorted((a for a in analises if a.esta_na_proplan), key=lambda a: -a.dias_unidade_atual)
    if not atuais:
        return "<p>Nenhum processo analisado está atualmente em unidade da PROPLAN.</p>"
    linhas = []
    for a in atuais:
        p = a.permanencias[-1]
        situacao = (
            '<span class="tag wn">aguardando recebimento</span>'
            if a.aguardando_recebimento else f"recebido em {_d(p.recebimento)}"
        )
        cls = ' class="tag al"' if a.alertas else ""
        linhas.append(
            f"<tr><td><a href='#p{_e(a.processo.numero)}'>{_e(a.processo.numero)}</a></td>"
            f"<td>{_e(a.processo.assunto)}</td><td>{_tag_setor(a.setor_atual)}</td>"
            f"<td>{_e(p.veio_de)}</td><td>{_d(p.inicio, True)}</td><td>{situacao}</td>"
            f'<td class="n" data-v="{a.dias_unidade_atual:.3f}"><span{cls}>{duracao(a.dias_unidade_atual)}</span></td>'
            f'<td class="n">{a.dias_uteis_unidade_atual}</td>'
            f'<td class="n" data-v="{a.dias_total_proplan:.3f}">{duracao(a.dias_total_proplan)}</td></tr>'
        )
    return (
        '<input type="search" placeholder="Filtrar…" data-filtra="#estoque tbody tr">'
        '<div class="tw"><table id="estoque"><thead><tr><th>Processo</th><th>Assunto</th><th>Setor</th>'
        '<th>Veio de</th><th>Chegou em</th><th>Situação</th><th class="n">Dias no setor</th>'
        '<th class="n">Dias úteis</th><th class="n">Total na PROPLAN</th></tr></thead><tbody>'
        + "".join(linhas) + "</tbody></table></div>"
    )


def _tabela_processos(analises: list[AnaliseProcesso], cfg: Config) -> str:
    siglas = cfg.siglas
    cab = "".join(f'<th class="n">{_e(s)}</th>' for s in siglas)
    linhas = []
    for a in sorted(analises, key=lambda a: -a.dias_total_proplan):
        cols = "".join(
            f'<td class="n" data-v="{a.dias_por_setor.get(s, 0):.3f}">'
            + (duracao(a.dias_por_setor[s]) + (f" <small>({a.passagens_por_setor[s]}x)</small>"
               if a.passagens_por_setor.get(s, 0) > 1 else "")
               if s in a.dias_por_setor else "—")
            + "</td>"
            for s in siglas
        )
        atual = _tag_setor(a.setor_atual) or _e(a.unidade_atual)
        if a.encerrado:
            atual += f' <span class="tag">{_e(a.processo.status)}</span>'
        linhas.append(
            f"<tr><td><a href='#p{_e(a.processo.numero)}'>{_e(a.processo.numero)}</a></td>"
            f"<td>{_e(a.processo.assunto)}</td><td>{atual}</td>"
            f'<td class="n" data-v="{a.dias_unidade_atual:.3f}">{duracao(a.dias_unidade_atual)}</td>'
            f"{cols}"
            f'<td class="n" data-v="{a.dias_total_proplan:.3f}"><b>{duracao(a.dias_total_proplan)}</b></td>'
            f'<td class="n" data-v="{a.dias_tramitacao_total:.3f}">{duracao(a.dias_tramitacao_total)}</td></tr>'
        )
    return (
        '<input type="search" placeholder="Filtrar por número, assunto, setor…" '
        'data-filtra="#todos tbody tr"><div class="tw"><table id="todos"><thead><tr>'
        '<th>Processo</th><th>Assunto</th><th>Local atual</th><th class="n">Dias no local atual</th>'
        f'{cab}<th class="n">Total PROPLAN</th><th class="n">Tramitação total</th></tr></thead><tbody>'
        + "".join(linhas) + "</tbody></table></div><p><small>Dias corridos acumulados em cada setor "
        "(soma de todas as passagens; entre parênteses, nº de passagens quando > 1).</small></p>"
    )


def _lista_contagem(titulo: str, contagem) -> str:
    if not contagem:
        return f"<div><h2>{_e(titulo)}</h2><p>—</p></div>"
    maior = max(contagem.values())
    linhas = "".join(
        f"<tr><td>{_e(k)}</td><td class='n'>{v}</td><td style='width:35%'>"
        f"<div class='bar' style='width:{100 * v / maior:.0f}%'></div></td></tr>"
        for k, v in contagem.most_common(15)
    )
    return (f"<div><h2>{_e(titulo)}</h2><div class='tw'><table><thead><tr><th>Unidade</th>"
            f"<th class='n'>Qtde</th><th></th></tr></thead><tbody>{linhas}</tbody></table></div></div>")


def _detalhe(a: AnaliseProcesso) -> str:
    p = a.processo
    linhas = []
    for perm in a.permanencias:
        receb = _d(perm.recebimento, True) if perm.recebimento else (
            '<span class="tag wn">não recebido</span>' if perm.aberta else "—"
        )
        linhas.append(
            f"<tr{' class=pp' if perm.setor else ''}><td>{_tag_setor(perm.setor)} {_e(perm.unidade)}</td>"
            f"<td>{_e(perm.veio_de)}</td><td>{_d(perm.inicio, True)}</td><td>{receb}</td>"
            f"<td>{_d(perm.fim, True) if perm.fim else '<b>atual</b>'}</td>"
            f'<td class="n">{duracao(perm.dias)}</td><td class="n">{perm.dias_uteis}</td>'
            f'<td class="n">{duracao(perm.dias_aguardando_recebimento)}</td></tr>'
        )
    alertas = "".join(f'<li><span class="tag al">alerta</span> {_e(x)}</li>' for x in a.alertas)
    interessados = "; ".join(p.interessados[:5])
    qtd_alertas = f' <span class="tag al">{len(a.alertas)} alerta(s)</span>' if a.alertas else ""
    return (
        f"<details id='p{_e(p.numero)}'><summary><b>{_e(p.numero)}</b> — {_e(p.assunto)} "
        f"{_tag_setor(a.setor_atual) if a.esta_na_proplan else ''}"
        f"{qtd_alertas}"
        f"</summary><p><small>Status: {_e(p.status or '—')} · Autuação: {_d(p.data_autuacao)} · "
        f"Origem: {_e(p.unidade_origem or '—')} · Natureza: {_e(p.natureza or '—')}"
        f"{' · Interessados: ' + _e(interessados) if interessados else ''}</small></p>"
        f"{('<p>' + _e(p.assunto_detalhado) + '</p>') if p.assunto_detalhado else ''}"
        f"{('<ul>' + alertas + '</ul>') if alertas else ''}"
        '<div class="tw"><table><thead><tr><th>Unidade</th><th>Veio de</th><th>Enviado em</th>'
        '<th>Recebido em</th><th>Saiu em</th><th class="n">Dias</th><th class="n">Dias úteis</th>'
        '<th class="n">Dias até receber</th></tr></thead><tbody>'
        + "".join(linhas) + "</tbody></table></div></details>"
    )


def gerar_html(c: Consolidado, cfg: Config, destino: Path, erros: list[str] | None = None) -> Path:
    analises = c.analises
    proplan = [a for a in analises if a.passou_pela_proplan]
    atuais = [a for a in analises if a.esta_na_proplan]
    com_alerta = [a for a in proplan if a.alertas]
    medio_total = (sum(a.dias_total_proplan for a in proplan) / len(proplan)) if proplan else 0

    alertas_html = "".join(
        f"<tr><td><a href='#p{_e(a.processo.numero)}'>{_e(a.processo.numero)}</a></td>"
        f"<td>{_tag_setor(a.setor_atual) or _e(a.unidade_atual)}</td>"
        f"<td>{'<br>'.join(_e(x) for x in a.alertas)}</td></tr>"
        for a in sorted(com_alerta, key=lambda a: -a.dias_unidade_atual)
    )
    erros_html = ""
    if erros:
        erros_html = ("<section><h2>Falhas de leitura</h2><ul>"
                      + "".join(f"<li>{_e(x)}</li>" for x in erros) + "</ul></section>")

    corpo = f"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Tramitação PROPLAN</title><style>{_CSS}</style></head><body>
<header><h1>Tramitação de processos na PROPLAN</h1>
<div class="sub">Fonte: SIPAC/UFPB (consulta pública) · Referência: {_d(c.referencia, True)} ·
Setores: {", ".join(_e(s) for s in cfg.siglas)}</div></header><main>
<div class="kpis">
<div class="kpi"><b>{len(analises)}</b><span>processos analisados</span></div>
<div class="kpi"><b>{len(proplan)}</b><span>passaram pela PROPLAN</span></div>
<div class="kpi"><b>{len(atuais)}</b><span>estão hoje na PROPLAN</span></div>
<div class="kpi"><b>{sum(1 for a in atuais if a.dias_unidade_atual >= cfg.dias_parado_no_setor)}</b>
<span>parados ≥ {cfg.dias_parado_no_setor} dias</span></div>
<div class="kpi"><b>{duracao(medio_total)}</b><span>tempo médio de cada processo na PROPLAN</span></div>
<div class="kpi"><b>{len(com_alerta)}</b><span>processos com alerta</span></div>
</div>
<section><h2>Desempenho por setor</h2>{_tabela_setores(c, cfg)}</section>
<section><h2>Estoque atual na PROPLAN (onde está cada processo e há quanto tempo)</h2>{_tabela_estoque(analises)}</section>
<section><h2>Alertas</h2>{('<div class="tw"><table><thead><tr><th>Processo</th><th>Local</th><th>Alertas</th></tr></thead><tbody>' + alertas_html + '</tbody></table></div>') if alertas_html else '<p>Nenhum alerta.</p>'}</section>
<section class="grid2">{_lista_contagem("De onde chegam os processos", c.entradas)}{_lista_contagem("Para onde saem da PROPLAN", c.saidas)}</section>
<section><h2>Todos os processos</h2>{_tabela_processos(analises, cfg)}</section>
<section><h2>Movimentação detalhada por processo</h2>
<input type="search" placeholder="Filtrar…" data-filtra="#detalhes > details">
<div id="detalhes">{"".join(_detalhe(a) for a in sorted(analises, key=lambda a: a.processo.numero))}</div></section>
{erros_html}
</main><script>{_JS}</script></body></html>"""
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(corpo, encoding="utf-8")
    return destino


# ===================================================================== Excel


def gerar_excel(c: Consolidado, cfg: Config, destino: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def aba(titulo, cabecalho, linhas, primeira=False):
        ws = wb.active if primeira else wb.create_sheet()
        ws.title = titulo
        ws.append(cabecalho)
        for c_ in ws[1]:
            c_.font = Font(bold=True, color="FFFFFF")
            c_.fill = PatternFill("solid", fgColor="1F5FAE")
        for linha in linhas:
            ws.append(list(linha))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(cabecalho, 1):
            largura = max([len(str(col))] + [len(str(l[i - 1] or "")) for l in linhas[:300]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, largura + 2), 60)
        return ws

    r = lambda x: round(x, 2)  # noqa: E731

    aba("Resumo por setor",
        ["Setor", "Nome", "Processos", "Passagens", "Média dias/passagem", "Mediana", "Máximo",
         "Média dias até receber", "Em estoque hoje", "Idade média estoque",
         f"Estoque >= {cfg.dias_parado_no_setor} dias"],
        [[s.sigla, s.nome, s.processos, s.passagens, r(s.media), r(s.mediana), r(s.maximo),
          r(s.media_aguardando), s.em_estoque, r(s.idade_media_estoque), s.estoque_acima_limite]
         for s in c.setores], primeira=True)

    siglas = cfg.siglas
    aba("Processos",
        ["Processo", "Assunto", "Status", "Autuação", "Unidade atual", "Setor atual",
         "Aguardando recebimento", "Dias no local atual", "Dias úteis no local atual"]
        + [f"Dias {s}" for s in siglas] + [f"Passagens {s}" for s in siglas]
        + ["Total dias PROPLAN", "Tramitação total (dias)", "Alertas"],
        [[a.processo.numero, a.processo.assunto, a.processo.status, a.processo.data_autuacao,
          a.unidade_atual, a.setor_atual or "", "SIM" if a.aguardando_recebimento else "NÃO",
          r(a.dias_unidade_atual), a.dias_uteis_unidade_atual]
         + [r(a.dias_por_setor.get(s, 0)) for s in siglas]
         + [a.passagens_por_setor.get(s, 0) for s in siglas]
         + [r(a.dias_total_proplan), r(a.dias_tramitacao_total), " | ".join(a.alertas)]
         for a in c.analises])

    aba("Permanências",
        ["Processo", "Assunto", "Seq", "Unidade", "Setor PROPLAN", "Veio de", "Foi para",
         "Enviado em", "Recebido em", "Saiu em", "Atual", "Dias", "Dias úteis", "Dias até receber"],
        [[a.processo.numero, a.processo.assunto, i + 1, p.unidade, p.setor or "", p.veio_de,
          p.foi_para, p.inicio, p.recebimento, p.fim, "SIM" if p.aberta else "", r(p.dias),
          p.dias_uteis, r(p.dias_aguardando_recebimento)]
         for a in c.analises for i, p in enumerate(a.permanencias)])

    aba("Movimentações",
        ["Processo", "Seq", "Data envio", "Unidade origem", "Unidade destino", "Enviado por",
         "Data recebimento", "Recebido por", "Urgente"],
        [[a.processo.numero, m.ordem + 1, m.data_envio, m.unidade_origem, m.unidade_destino,
          m.enviado_por, m.data_recebimento, m.recebido_por, "SIM" if m.urgente else ""]
         for a in c.analises for m in a.processo.movimentacoes])

    aba("Fluxos", ["Origem", "Destino", "Quantidade"],
        [[o, d, q] for (o, d), q in c.fluxos.most_common()])

    for ws in wb.worksheets:
        for linha in ws.iter_rows(min_row=2):
            for cel in linha:
                if isinstance(cel.value, datetime):
                    cel.number_format = "dd/mm/yyyy hh:mm"

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


# ================================================================== terminal


def resumo_texto(a: AnaliseProcesso) -> str:
    p = a.processo
    out = [
        f"Processo {p.numero} — {p.assunto}",
        f"  Status: {p.status or '—'} | Autuação: {_d(p.data_autuacao)} | Origem: {p.unidade_origem or '—'}",
        f"  Local atual: {a.unidade_atual} [{a.setor_atual or 'fora da PROPLAN'}] há "
        f"{duracao(a.dias_unidade_atual)} ({a.dias_uteis_unidade_atual} {"dia útil" if a.dias_uteis_unidade_atual == 1 else "dias úteis"})"
        + (" — AGUARDANDO RECEBIMENTO" if a.aguardando_recebimento else ""),
        f"  Tramitação total: {duracao(a.dias_tramitacao_total)} | Na PROPLAN: {duracao(a.dias_total_proplan)}",
    ]
    if a.dias_por_setor:
        out.append("  Por setor: " + "; ".join(
            f"{s} {duracao(d)} ({a.passagens_por_setor[s]}x)" for s, d in a.dias_por_setor.items()))
    out.append("  Movimentação:")
    for perm in a.permanencias:
        marca = f"[{perm.setor}]" if perm.setor else "        "
        out.append(
            f"    {_d(perm.inicio, True):<16} {marca:<12} {perm.unidade[:60]:<60} "
            f"{duracao(perm.dias):>12}" + ("  <- atual" if perm.aberta else "")
        )
    for al in a.alertas:
        out.append(f"  ! {al}")
    return "\n".join(out)
